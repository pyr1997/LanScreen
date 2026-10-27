"""Integration checks, retaining every test artifact in LANSCREEN_TEST_ROOT."""
from pathlib import Path
import io
import json
import os
import subprocess
import sys
import time
import uuid
import zipfile

import fitz
from PIL import Image
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from binding_helpers import approved_headers
from app import create_app, find_ffmpeg, find_office

TEST_ROOT = Path(os.environ.get('LANSCREEN_TEST_ROOT', Path.cwd() / 'work' / 'api-tests'))


@pytest.fixture
def server():
    directory = TEST_ROOT / uuid.uuid4().hex
    app = create_app(directory, testing=True)
    client = app.test_client()
    setup = client.post('/api/setup', json={'setupCode': app.extensions['lanscreen']['setting']('setup_code'), 'password': 'TestOnly-Password-2026'})
    assert setup.status_code == 200
    csrf = {'X-CSRF-Token': setup.json['csrf']}
    yield app, client, csrf
    app.extensions['lanscreen']['executor'].shutdown(wait=True)


def upload_wait(server, name, content):
    app, client, csrf = server
    result = client.post('/api/media', data={'file': (io.BytesIO(content), name)}, headers=csrf)
    assert result.status_code == 202, result.json
    app.extensions['lanscreen']['executor'].submit(lambda: None).result(timeout=210)
    return next(x for x in client.get('/api/dashboard').json['media'] if x['id'] == result.json['id'])


def image_bytes():
    out = io.BytesIO()
    Image.new('RGB', (400, 225), (45, 98, 244)).save(out, 'PNG')
    return out.getvalue()


def test_auth_origin_limits(server):
    app, client, csrf = server
    stranger = app.test_client()
    assert stranger.get('/api/dashboard').status_code == 401
    assert stranger.get('/api/player/manifest').status_code == 401
    assert stranger.post('/api/setup', json={'setupCode':'bogus','password':'new-password!'}).status_code == 409
    assert client.put('/api/playlist', json={'items':[]}).status_code == 403
    assert client.post('/api/publish', json={}, headers={**csrf, 'Origin':'https://attacker.example'}).status_code == 403
    assert stranger.post('/api/login', json={'password':'x'*9000}).status_code == 413
    assert stranger.post('/api/login', json={'password':'bad'}).status_code == 401
    assert stranger.post('/api/login', json={'password':'TestOnly-Password-2026'}).status_code == 200


@pytest.mark.parametrize('origin', ['http://[', 'http://]', 'http://[invalid]'])
def test_malformed_origin_is_forbidden_without_mutation_or_error_log(server, caplog, origin):
    app, client, csrf = server
    previous = client.get('/api/dashboard').json
    headers = {**csrf, 'Origin': origin}
    response = client.put('/api/display-token',
                          json={'displayToken': 'Next.Playback_2026', 'currentDisplayToken': previous['displayToken']},
                          headers=headers)
    assert response.status_code == 403
    assert client.post('/api/publish', json={'items': []}, headers=headers).status_code == 403
    current = client.get('/api/dashboard').json
    for key in ('media', 'devices', 'draft', 'published', 'revision', 'cachePolicy', 'displayToken'):
        assert current[key] == previous[key]
    assert not any(record.levelno >= 40 for record in caplog.records)


def test_same_origin_api_update_remains_allowed(server):
    app, client, csrf = server
    token = client.get('/api/dashboard').json['displayToken']
    response = client.put('/api/display-token', json={'displayToken': token, 'currentDisplayToken': token},
                          headers={**csrf, 'Origin': 'http://localhost'})
    assert response.status_code == 200 and response.json == {'ok': True, 'changed': False}


@pytest.mark.parametrize('authorization', ['Bearer é', 'Bearer ' + 'a' * 129])
def test_malformed_player_credentials_are_unauthorized_without_mutation(server, caplog, authorization):
    app, client, csrf = server
    previous = client.get('/api/dashboard').json
    player = app.test_client()
    headers = {'Authorization': authorization}
    assert player.get('/api/player/manifest', headers=headers).status_code == 401
    assert player.get('/media/unknown/page.jpg', headers=headers).status_code == 401
    assert player.post('/api/player/heartbeat', json={'deviceId': 'unauthorized-tv'}, headers=headers).status_code == 401
    current = client.get('/api/dashboard').json
    for key in ('media', 'devices', 'draft', 'published', 'revision', 'cachePolicy', 'displayToken'):
        assert current[key] == previous[key]
    assert not any(record.levelno >= 40 for record in caplog.records)


@pytest.mark.parametrize('csrf_token', ['é', 'a' * 49])
def test_malformed_csrf_is_forbidden_without_mutation(server, caplog, csrf_token):
    app, client, csrf = server
    original = client.get('/api/dashboard').json['displayToken']
    response = client.put('/api/display-token',
                          json={'displayToken': 'Next.Playback_2026', 'currentDisplayToken': original},
                          headers={'X-CSRF-Token': csrf_token})
    assert response.status_code == 403
    assert client.get('/api/dashboard').json['displayToken'] == original
    assert not any(record.levelno >= 40 for record in caplog.records)


@pytest.mark.parametrize('setup_code', ['é', '中文初始化码', 'a' * 9, None, False, 123, [], {}])
def test_malformed_setup_code_is_forbidden_without_initialization(caplog, setup_code):
    application = create_app(TEST_ROOT / uuid.uuid4().hex, testing=True)
    try:
        client = application.test_client()
        setting = application.extensions['lanscreen']['setting']
        original = setting('setup_code')
        response = client.post('/api/setup', json={'setupCode': setup_code, 'password': 'TestOnly-Password-2026'})
        assert response.status_code == 403
        assert setting('password') is None and setting('setup_code') == original
        assert client.post('/api/setup', json={'setupCode': original, 'password': 'TestOnly-Password-2026'}).status_code == 200
        assert not any(record.levelno >= 40 for record in caplog.records)
    finally:
        application.extensions['lanscreen']['executor'].shutdown(wait=True)


def test_display_token_requires_admin_csrf_and_same_origin(server):
    app, client, csrf = server
    original = client.get('/api/dashboard').json['displayToken']
    body = {'displayToken': 'LanScreen_Test-2026', 'currentDisplayToken': original}
    stranger = app.test_client()
    assert stranger.put('/api/display-token', json=body, headers=csrf).status_code == 401
    assert client.put('/api/display-token', json=body).status_code == 403
    assert client.put('/api/display-token', json=body, headers={'X-CSRF-Token': 'wrong'}).status_code == 403
    assert client.put('/api/display-token', json=body,
                      headers={**csrf, 'Origin': 'https://attacker.example'}).status_code == 403
    assert client.get('/api/dashboard').json['displayToken'] == original


@pytest.mark.parametrize('value', [
    None, False, 123456789012, [], {}, '', 'a' * 11, 'a' * 129,
    ' ValidToken-2026', 'ValidToken-2026 ', 'Valid Token-2026',
    'ValidToken-2026\n', 'ValidToken-2026\r', 'ValidToken\t2026',
    'ValidToken/2026', 'ValidToken+2026', 'ValidToken=2026',
    'ValidToken:2026', 'ValidToken中文2026', 'ValidTokené2026',
])
def test_display_token_rejects_invalid_values_without_changing_key(server, value):
    app, client, csrf = server
    original = client.get('/api/dashboard').json['displayToken']
    response = client.put('/api/display-token', json={'displayToken': value, 'currentDisplayToken': original}, headers=csrf)
    assert response.status_code == 400
    assert client.get('/api/dashboard').json['displayToken'] == original
    assert app.test_client().get('/api/player/manifest',
                                 headers=approved_headers(server)).status_code == 200


@pytest.mark.parametrize('token', ['Ab12._~-xyZ9', 'Ab12._~-' * 16])
def test_display_token_accepts_length_boundaries_and_is_idempotent(server, token):
    app, client, csrf = server
    original = client.get('/api/dashboard').json['displayToken']
    response = client.put('/api/display-token', json={'displayToken': token, 'currentDisplayToken': original}, headers=csrf)
    assert response.status_code == 200 and response.json == {'ok': True, 'changed': True}
    assert client.get('/api/dashboard').json['displayToken'] == token
    repeat = client.put('/api/display-token', json={'displayToken': token, 'currentDisplayToken': token}, headers=csrf)
    assert repeat.status_code == 200 and repeat.json == {'ok': True, 'changed': False}
    assert app.test_client().get('/api/player/manifest',
                                 headers=approved_headers(server)).status_code == 200


@pytest.mark.parametrize('body', [
    {}, {'currentDisplayToken': None}, {'currentDisplayToken': False},
    {'currentDisplayToken': 123}, {'currentDisplayToken': []}, {'currentDisplayToken': {}},
])
def test_display_token_requires_current_key_string(server, body):
    app, client, csrf = server
    original = client.get('/api/dashboard').json['displayToken']
    response = client.put('/api/display-token', json={**body, 'displayToken': 'Next.Key_2026'}, headers=csrf)
    assert response.status_code == 400
    assert client.get('/api/dashboard').json['displayToken'] == original


@pytest.mark.parametrize('expected', ['', 'Incorrect_Key_2026', '当前播放密钥', '\ud800'])
def test_display_token_rejects_incorrect_current_key_without_error_or_mutation(server, expected):
    app, client, csrf = server
    original = client.get('/api/dashboard').json['displayToken']
    response = client.put('/api/display-token',
                          json={'displayToken': 'Next.Key_2026', 'currentDisplayToken': expected}, headers=csrf)
    assert response.status_code == 409
    assert client.get('/api/dashboard').json['displayToken'] == original


def test_display_token_stale_admin_page_cannot_overwrite_another_change(server):
    app, client_a, csrf_a = server
    client_b = app.test_client()
    login_b = client_b.post('/api/login', json={'password': 'TestOnly-Password-2026'})
    assert login_b.status_code == 200
    csrf_b = {'X-CSRF-Token': login_b.json['csrf']}
    snapshot_a = client_a.get('/api/dashboard').json['displayToken']
    snapshot_b = client_b.get('/api/dashboard').json['displayToken']
    assert snapshot_a == snapshot_b
    token_a, token_b = 'Page-A.Playback_2026', 'Page-B.Playback_2026'
    first = client_a.put('/api/display-token',
                         json={'displayToken': token_a, 'currentDisplayToken': snapshot_a}, headers=csrf_a)
    assert first.status_code == 200 and first.json == {'ok': True, 'changed': True}
    conflict = client_b.put('/api/display-token',
                            json={'displayToken': token_b, 'currentDisplayToken': snapshot_b}, headers=csrf_b)
    assert conflict.status_code == 409
    assert not any(secret in conflict.get_data(as_text=True) for secret in (snapshot_a, token_a, token_b))
    stale_repeat = client_b.put('/api/display-token',
                                json={'displayToken': token_a, 'currentDisplayToken': snapshot_b}, headers=csrf_b)
    assert stale_repeat.status_code == 409
    assert client_a.get('/api/dashboard').json['displayToken'] == token_a
    player = app.test_client()
    bound_headers = approved_headers(server)
    for token, status in ((snapshot_a, 401), (token_b, 401), (token_a, 200)):
        assert player.get('/api/player/manifest', headers={**bound_headers, 'Authorization': 'Bearer ' + token}).status_code == status


def test_display_token_rotation_revokes_player_access_and_survives_restart(server):
    app, client, csrf = server
    media = upload_wait(server, 'rotation-poster.png', image_bytes())
    assert media['status'] == 'ready'
    entries = [{'mediaId': media['id'], 'duration': 3}]
    assert client.post('/api/publish', json={'items': entries}, headers=csrf).status_code == 200
    previous = client.get('/api/dashboard').json
    old_bearer = approved_headers(server, 'token-rotation-tv')
    new_token = 'New_Playback.Key~2026'
    new_bearer = {**old_bearer, 'Authorization': 'Bearer ' + new_token}
    player = app.test_client()
    manifest = player.get('/api/player/manifest', headers=old_bearer).json
    media_url = manifest['items'][0]['url']
    assert client.put('/api/display-token',
                      json={'displayToken': new_token, 'currentDisplayToken': previous['displayToken']},
                      headers=csrf).json['changed'] is True
    assert player.get('/api/player/manifest', headers=old_bearer).status_code == 401
    assert player.get(media_url, headers=old_bearer).status_code == 401
    heartbeat = {'protocolVersion': 3, 'deviceId': 'token-rotation-tv', 'name': 'Test TV', 'status': 'playing'}
    assert player.post('/api/player/heartbeat', json=heartbeat, headers=old_bearer).status_code == 401
    current_manifest = player.get('/api/player/manifest', headers=new_bearer)
    assert current_manifest.status_code == 200
    assert current_manifest.json['items'] == manifest['items']
    assert current_manifest.json['revision'] == manifest['revision']
    with player.get(media_url, headers=new_bearer) as response:
        assert response.status_code == 200 and response.data
    assert player.post('/api/player/heartbeat', json=heartbeat, headers=new_bearer).status_code == 200
    assert player.put('/api/display-token', json={'displayToken': 'NotAnAdmin-2026'},
                      headers=new_bearer).status_code == 401
    assert player.get('/api/dashboard', headers=new_bearer).status_code == 401
    current = client.get('/api/dashboard').json
    for key in ('media', 'draft', 'published', 'revision', 'cachePolicy'):
        assert current[key] == previous[key]

    app.extensions['lanscreen']['executor'].shutdown(wait=True)
    restarted = create_app(app.extensions['lanscreen']['data'], testing=True)
    try:
        restarted_client = restarted.test_client()
        assert restarted_client.get('/api/player/manifest', headers=old_bearer).status_code == 401
        assert restarted_client.get('/api/player/manifest', headers=new_bearer).json['items'] == manifest['items']
        assert restarted_client.post('/api/login', json={'password': 'TestOnly-Password-2026'}).status_code == 200
        assert restarted_client.get('/api/dashboard').json['displayToken'] == new_token
    finally:
        restarted.extensions['lanscreen']['executor'].shutdown(wait=True)


def test_publish_image_and_protected_range(server):
    app, client, csrf = server
    media = upload_wait(server, '../../海报<script>.png', image_bytes())
    assert media['status'] == 'ready'
    assert media['name'] == '海报<script>.png'
    token = client.get('/api/dashboard').json['displayToken']
    player = app.test_client()
    bearer = approved_headers(server)
    entries = [{'mediaId':media['id'], 'duration':3}]
    assert client.put('/api/playlist', json={'items':entries}, headers=csrf).status_code == 200
    assert player.get('/api/player/manifest', headers=bearer).json['items'] == []
    # Atomic publish takes caller's own snapshot, not another browser's stored draft.
    assert client.post('/api/publish', json={'items':entries}, headers=csrf).status_code == 200
    manifest = player.get('/api/player/manifest', headers=bearer).json
    assert manifest['revision'] == 1 and manifest['items'][0]['duration'] == 3
    url = manifest['items'][0]['url']
    assert player.get(url).status_code == 401
    response = player.get(url, headers=bearer)
    assert response.status_code == 200 and response.headers['Cache-Control'].startswith('private')
    assert 'Authorization' in response.headers['Vary']
    import hashlib
    assert hashlib.sha256(response.data).hexdigest() == manifest['items'][0]['sha256']
    assert player.get(url, headers={**bearer, 'Range':'bytes=0-15'}).status_code == 206
    assert player.get('/media/'+media['id']+'/original.png', headers=bearer).status_code == 404
    assert client.post('/api/media/'+media['id']+'/archive', headers=csrf).status_code == 409
    assert player.post('/api/player/heartbeat',json={'protocolVersion':3,'deviceId':'tv-test','name':'大厅电视','status':'playing'},headers=bearer).status_code == 200
    assert client.get('/api/dashboard').json['devices'][0]['online'] is True
    assert client.post('/api/publish',json={'items':[]},headers=csrf).status_code==200
    assert player.get('/api/player/manifest',headers=bearer).json['items']==[]
    assert client.post('/api/media/'+media['id']+'/archive',headers=csrf).status_code==200
    assert (app.extensions['lanscreen']['data']/'media'/media['id']/'original.png').exists()


def test_documents_and_invalid_inputs(server):
    doc=fitz.open()
    for label in ('First page', 'Second page'):
        page=doc.new_page(width=960,height=540)
        page.insert_text((60,80),label,fontsize=30)
    media=upload_wait(server,'company.pdf',doc.tobytes())
    assert media['status']=='ready' and len(media['pages'])==2
    assert upload_wait(server,'broken.pdf',b'not a pdf')['status']=='error'
    assert upload_wait(server,'broken.png',b'not an image')['status']=='error'
    app,client,csrf=server
    assert client.post('/api/media',data={'file':(io.BytesIO(b'<script>'),'bad.html')},headers=csrf).status_code==400
    assert client.post('/api/media',data={'file':(io.BytesIO(b'bad'),'bad.pptx')},headers=csrf).status_code==400
    assert client.put('/api/playlist',json={'items':[{'mediaId':media['id'],'duration':0}]},headers=csrf).status_code==400
    assert client.put('/api/playlist',json={'items':[{'mediaId':'missing','duration':5}]},headers=csrf).status_code==400


def test_real_video_conversion_and_fake_mp4(server):
    ffmpeg=find_ffmpeg()
    assert ffmpeg, 'Video test requires installed imageio-ffmpeg or ffmpeg'
    source=server[0].extensions['lanscreen']['data']/'source-video.mp4'
    subprocess.run([ffmpeg,'-v','error','-f','lavfi','-i','color=c=0x2d62f4:s=640x360:d=1','-c:v','libx264','-pix_fmt','yuv420p',str(source)],check=True,timeout=30)
    media=upload_wait(server,'demo.mp4',source.read_bytes())
    assert media['status']=='ready' and media['pages'][0]['type']=='video'
    assert upload_wait(server,'fake.mp4',b'ftypnot-a-video')['status']=='error'


@pytest.mark.skipif(not find_office(),reason='LibreOffice not installed')
def test_pptx_conversion(server):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    presentation=Presentation()
    presentation.slide_width=Inches(16)
    presentation.slide_height=Inches(9)
    for label in ('公司宣传 · 中文验证','局域网轮播 · 第二页'):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1),Inches(2),Inches(14),Inches(2))
        shape.text_frame.paragraphs[0].text=label
        shape.text_frame.paragraphs[0].font.name='Arial Unicode MS'
        shape.text_frame.paragraphs[0].font.size=Pt(40)
    stream=io.BytesIO()
    presentation.save(stream)
    media=upload_wait(server,'公司介绍.pptx',stream.getvalue())
    assert media['status']=='ready',media['error']
    assert len(media['pages'])==2


def test_expanded_page_limit(server):
    app,client,csrf=server
    media=upload_wait(server,'test.png',image_bytes())
    with app.extensions['lanscreen']['db']() as c:
        pages=media['pages']*200
        c.execute('UPDATE media SET pages=? WHERE id=?',(json.dumps(pages),media['id']))
    response=client.post('/api/publish',json={'items':[{'mediaId':media['id'],'duration':10}]*11},headers=csrf)
    assert response.status_code==400 and '2000' in response.json['error']
